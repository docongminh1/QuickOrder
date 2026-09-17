# -*- coding: utf-8 -*-
"""Thêm thuốc đọc từ ảnh tủ vào file Excel của quầy (không đụng file mẫu).
Chạy: python3 tools/add_from_photo.py  → CatLieuNhanh_quay.xlsx
"""
import os, shutil
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Side, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'CatLieuNhanh_mau.xlsx')
DST = os.path.join(ROOT, 'CatLieuNhanh_quay.xlsx')
TU = 'Tủ kê đơn'
NOTE = 'đọc từ ảnh 17/09, dược sĩ rà lại'

# Tên thuốc | Hoạt chất | mg | Dạng | Hãng | Đối tượng | Còn hàng | Kê đơn | Vị trí | Ghi chú
ROWS = [
 # hàng 1
 ('Auclanityl (?)', 'Amoxicillin + Acid clavulanic', '?', 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 1 · trái', 'hàm lượng chưa đọc được (875/125?) — ' + NOTE),
 ('Auclanityl 500/125', 'Amoxicillin + Acid clavulanic', '500 + 125', 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 1 · giữa', NOTE),
 ('Amoxicillin 500', 'Amoxicillin', 500, 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 1 · giữa', NOTE),
 ('Ampicillin MKP 500', 'Ampicillin', 500, 'viên', 'Mekophar', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 1 · phải', NOTE),
 # hàng 2
 ('Vipocef 100', 'Cefpodoxim', 100, 'viên', '', 'Cả hai', 'Có', 'Có', f'{TU} · hàng 2 · trái', NOTE),
 ('Vipocef 200', 'Cefpodoxim', 200, 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 2 · phải', NOTE),
 # hàng 3
 ('Cefixim 200', 'Cefixim', 200, 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 3 · trái', NOTE),
 ('Doxycyclin 100 (?)', 'Doxycyclin', 100, 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 3 · trái', 'hàm lượng đoán 100 — ' + NOTE),
 ('Cephalexin 500', 'Cephalexin', 500, 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 3 · trái', NOTE),
 ('Cotrimxazon 960', 'Sulfamethoxazol + Trimethoprim', '800 + 160', 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 3 · giữa', NOTE),
 ('Hộp giữa hàng 3 (?)', '?', '', 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 3 · giữa', '2 hộp trắng/xanh nhạt chưa đọc được tên — dược sĩ điền'),
 ('Agitro 500', 'Azithromycin', 500, 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 3 · phải', NOTE),
 ('Zolmed 150', 'Fluconazol', 150, 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 3 · phải', NOTE),
 ('Ifatrax (?)', '?', '', 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 3 · phải', 'chưa chắc hoạt chất — ' + NOTE),
 # hàng 4
 ('Cataflam 25', 'Diclofenac', 25, 'viên', 'Novartis', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 4 · trái', NOTE),
 ('Cataflam 50', 'Diclofenac', 50, 'viên', 'Novartis', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 4 · trái', NOTE),
 ('Diclofenac 75 (?)', 'Diclofenac', 75, 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 4 · trái', 'hộp nhỏ, chưa chắc — ' + NOTE),
 ('Celecoxib 200 (?)', 'Celecoxib', 200, 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 4 · giữa', 'hàm lượng đoán 200 — ' + NOTE),
 ('Kamelox 15', 'Meloxicam', 15, 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 4 · giữa', NOTE),
 ('Meloxicam 7,5', 'Meloxicam', 7.5, 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 4 · giữa', NOTE),
 ('Mydocalm 50', 'Tolperison', 50, 'viên', 'Gedeon Richter', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 4 · giữa', NOTE),
 ('Respamax (?)', '?', '', 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 4 · phải', 'tên đọc mờ — ' + NOTE),
 ('Bitolrison 150 (?)', '?', 150, 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 4 · phải', 'tên đọc mờ — ' + NOTE),
 # hàng 5
 ('Cytan (?)', '?', '', 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 5 · trái', 'hộp đỏ, tên đọc mờ — ' + NOTE),
 ('Alpha Choay', 'Alphachymotrypsin', 4.2, 'viên', 'Sanofi', 'Người lớn', 'Có', '', f'{TU} · hàng 5 · trái', NOTE),
 ('Medrol 4 (?)', 'Methylprednisolon', 4, 'viên', 'Pfizer', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 5 · giữa', 'có 2 loại hộp, đoán 4 và 16 — ' + NOTE),
 ('Medrol 16 (?)', 'Methylprednisolon', 16, 'viên', 'Pfizer', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 5 · giữa', 'có 2 loại hộp, đoán 4 và 16 — ' + NOTE),
 ('Bambuterol 10 A.T', 'Bambuterol', 10, 'viên', 'An Thiên', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 5 · giữa', NOTE),
 ('Neo-Codion', 'Codein + Sulfogaiacol + Cao mềm Grindelia', '', 'viên', 'Bouchara', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 5 · phải', NOTE),
 # hàng 6
 ('Cidetuss (?)', '?', '', 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 6 · giữa', 'chưa chắc hoạt chất — ' + NOTE),
 ('Star+ (?)', '?', '', 'viên', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 6 · giữa', '4 hộp nâu/xanh, chưa đọc được — ' + NOTE),
 ('Hộp trái hàng 6 (?)', '?', '', 'hộp', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 6 · trái', 'vài hộp xanh/vàng chưa đọc được — dược sĩ điền'),
 ('Hộp phải hàng 6 (?)', '?', '', 'hộp', '', 'Người lớn', 'Có', 'Có', f'{TU} · hàng 6 · phải', 'hộp nhỏ chưa đọc được — dược sĩ điền'),
]

shutil.copyfile(SRC, DST)
wb = load_workbook(DST)
ws = wb['Thuốc']
first = ws.max_row + 1
thin = Side(style='thin', color='D3DDD8'); border = Border(left=thin, right=thin, top=thin, bottom=thin)
mark = PatternFill('solid', fgColor='FFF4CC')  # vàng nhạt: dòng mới từ ảnh
for r in ROWS:
    ws.append(list(r))
    for c in ws[ws.max_row]:
        c.border = border; c.alignment = Alignment(vertical='center', wrap_text=True)
        if '(?)' in str(r[0]) or r[1] == '?': c.fill = mark
    ws.row_dimensions[ws.max_row].height = 22
ws.auto_filter.ref = f"A1:J{ws.max_row}"
# ghi chú ở sheet Hướng dẫn
hd = wb['Hướng dẫn']
hd.cell(row=hd.max_row + 2, column=1, value=f'File này là bản của QUẦY (không phải mẫu): đã thêm {len(ROWS)} thuốc đọc từ ảnh Tủ kê đơn (dòng {first}–{ws.max_row} sheet Thuốc). Dòng tô vàng có "(?)" là chưa chắc, dược sĩ sửa tên/hoạt chất/mg rồi bỏ dấu (?).')
wb.save(DST)
print(f'{DST}: thêm {len(ROWS)} thuốc, dòng {first}–{ws.max_row}')
